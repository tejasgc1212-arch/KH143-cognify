"""Extract uploaded text in memory. Uploaded filenames are never disk paths."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TOTAL_BYTES = 50 * 1024 * 1024
MAX_TEXT_CHARS = 100_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class DocumentError(ValueError):
    """An upload needs to be corrected before it can be analyzed."""


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentError("Use a PDF, DOCX, or TXT file. Convert legacy DOC files to DOCX.")
    if not content:
        raise DocumentError("The uploaded file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise DocumentError("Each file must be 10 MB or smaller.")

    try:
        if extension == ".txt":
            encoding = "utf-16" if content.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            text = content.decode(encoding)
            if "\x00" in text:
                raise DocumentError("The TXT file must contain readable text.")
        elif extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted and not reader.decrypt(""):
                raise DocumentError("Remove the PDF password before uploading.")
            if len(reader.pages) > 100:
                raise DocumentError("PDF files may contain at most 100 pages.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            # DOCX is a ZIP archive. Bound the expanded size before parsing XML.
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_TOTAL_BYTES:
                    raise DocumentError("The expanded DOCX file is too large.")
            document = Document(BytesIO(content))
            chunks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                chunks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            for section in document.sections:
                chunks.extend(paragraph.text for paragraph in section.header.paragraphs)
                chunks.extend(paragraph.text for paragraph in section.footer.paragraphs)
            text = "\n".join(chunks)
    except DocumentError:
        raise
    except UnicodeError:
        raise DocumentError("Save the TXT file as UTF-8 or UTF-16 and upload it again.") from None
    except Exception:
        # Parser exceptions can contain file content or local paths.
        raise DocumentError("The document could not be read. Re-export it as PDF, DOCX, or TXT.") from None

    text = text.strip()
    if not text:
        raise DocumentError("No readable text was found. Scanned PDFs need OCR before uploading.")
    if len(text) > MAX_TEXT_CHARS:
        raise DocumentError("The extracted text exceeds 100,000 characters. Upload a shorter document.")
    return text
