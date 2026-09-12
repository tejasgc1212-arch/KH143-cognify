"""Integration checks with synthetic documents. No provider credentials required."""

from collections import Counter
from io import BytesIO
import json

from docx import Document
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest

from backend import orchestrator
from backend.main import app
from backend.models import GitHubProfile
from backend.parsers import DocumentError, extract_text
from backend.utils import AgentOutputError, normalize_agent_output, score_value


class FakeAgents:
    """Exercise real routing/adapters while replacing only remote agent calls."""

    def __init__(self):
        self.calls = []

    def invoke(self, stage, payload):
        self.calls.append((stage, payload))
        if stage == "job":
            return AIMessage(content='```json\n{"required_skills":["Python"],"required_experience":"2 years","education":"Degree","responsibilities":["Build APIs"]}\n```')
        if stage == "resume":
            if "BROKEN_AGENT" in payload["resume"]:
                raise RuntimeError("private provider diagnostic")
            name = payload["resume"].splitlines()[0]
            return {"name": name, "skills": ["Python", "SQL"], "education": "CS degree",
                    "experience": "3 years building APIs", "projects": ["Resume-only API project"]}
        if stage == "github":
            return AIMessage(content="Project purpose: Example API\nTechnologies used: Python, FastAPI\nProject Score: 85/100")
        if stage == "matching":
            name = json.loads(payload["resume_data"])["name"]
            score = 91 if name == "Alice Example" else 73
            return f"Candidate Match Score: {score}/100\nRequired Skills: 36/40\nGitHub Evidence: 15/25\nExperience: 16/20\nEducation: 10/10\nAdditional Skills: 4/5\nMatched Skills:\n- Python\nMissing Skills:\n- SQL\nFinal Recommendation:\nGood Match\nReason:\nRelevant API experience."
        if stage == "skill_gap":
            assert "candidate_match_score" in json.loads(payload["match_data"])
            return "Critical Skill Gaps:\n- SQL query optimization\nRecommended Improvements:\n- Practice SQL\nOverall Gap Level:\nMedium"
        if stage == "summary":
            resume = json.loads(payload["resume_data"])
            assert "critical_skill_gaps" in json.loads(payload["skill_gap_data"])
            return {"candidate": resume["name"], "recruiter_summary": resume["name"] + " has relevant API experience.",
                    "key_strengths": ["Python APIs"], "important_gaps": ["SQL optimization"],
                    "final_recommendation": "Recommend", "match_score": 1}
        raise AssertionError(stage)

    def github_data(self, url):
        self.calls.append(("github_fetch", {"url": url}))
        return "Repository Name: example\nLanguages: Python\nREADME: Example API"


@pytest.fixture
def agents(monkeypatch):
    fake = FakeAgents()
    monkeypatch.setattr(orchestrator, "get_agents", lambda: fake)
    return fake


@pytest.fixture
def client(agents):
    with TestClient(app, raise_server_exceptions=False) as client:
        yield client


def submit(client, *, data=None, files=None):
    fields = {"jobRole": "Python Developer", "department": "Engineering", "minimumExperience": "2",
              "location": "Remote", "employmentType": "Full-time", "jobDescription": "Build Python APIs.",
              "githubProfiles": "[]"}
    fields.update(data or {})
    return client.post("/api/analyze", data=fields,
                       files=files if files is not None else [("resumes", ("alice.txt", b"Alice Example\nPython developer", "text/plain"))])


def pdf_bytes():
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 250 Td (Alice Example - Python Developer) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes():
    document = Document()
    document.add_paragraph("Alice Example")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "FastAPI"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize("output", [
    {"name": "Alice"}, '{"name":"Alice"}', '```json\n{"name":"Alice"}\n```',
    'Here is the result: {"name":"Alice"}', AIMessage(content='{"name":"Alice"}'),
    AIMessage(content=[{"type": "text", "text": '{"name":"Alice"}'}]),
    '**Name:** Alice', '## Name\nAlice',
])
def test_agent_output_formats(output):
    assert normalize_agent_output(output)["name"] == "Alice"


def test_normalization_retains_prose_and_never_executes():
    malicious = "__import__('os').system('do not execute')"
    assert normalize_agent_output(malicious)["raw_text"] == malicious
    assert normalize_agent_output('```json\n{"broken":\n```')["raw_text"]
    with pytest.raises(AgentOutputError):
        normalize_agent_output(None)
    assert score_value("91/100") == 91
    assert score_value("36/40", 40) == 36
    for invalid in [True, "Not Mentioned", "unknown 91", "101/100", "-1", "NaN", "70/0"]:
        assert score_value(invalid) is None


def test_document_parsers():
    assert "Python Developer" in extract_text("alice.pdf", pdf_bytes())
    assert "Python | FastAPI" in extract_text("alice.docx", docx_bytes())
    assert extract_text("unicode.txt", "Alice – Python".encode("utf-16")) == "Alice – Python"
    for filename, content in [("empty.txt", b""), ("old.doc", b"document"), ("broken.pdf", b"invalid"),
                              ("binary.txt", b"\x00data"), ("broken.docx", b"invalid")]:
        with pytest.raises(DocumentError):
            extract_text(filename, content)


def test_health_and_frontend_do_not_call_agents(client, agents):
    assert client.get("/api/health").json() == {"status": "ok"}
    assert client.get("/").status_code == 200
    assert "CognifyAPI" in client.get("/frontend/api.js").text
    assert client.get("/frontend/../Agents/.env").status_code == 404
    assert agents.calls == []


def test_multiple_candidates_are_isolated_and_sorted(client, agents):
    response = submit(client, files=[("resumes", ("bob.txt", b"Bob Example\nPython", "text/plain")),
                                     ("resumes", ("alice.txt", b"Alice Example\nPython", "text/plain"))])
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "completed"
    assert [(item["name"], item["matchScore"], item["rank"]) for item in result["candidates"]] == [
        ("Alice Example", 91, 1), ("Bob Example", 73, 2)]
    assert len({item["id"] for item in result["candidates"]}) == 2
    assert Counter(stage for stage, _ in agents.calls) == {"job": 1, "resume": 2, "matching": 2, "skill_gap": 2, "summary": 2}
    assert all(item["github"] == "Not Provided" and item["projects"] == [] for item in result["candidates"])
    assert result["candidates"][0]["skillsMatch"] == 90
    # Summary's deliberately conflicting score cannot overwrite Matching Agent.
    assert result["candidates"][0]["matchScore"] != 1


def test_jd_file_and_pasted_text_are_combined_once(client, agents):
    response = submit(client, files=[("jobDescriptionFile", ("job.txt", b"SQL preferred", "text/plain")),
                                     ("resumes", ("alice.txt", b"Alice Example", "text/plain"))])
    assert response.status_code == 200
    job_payloads = [payload["job_description"] for stage, payload in agents.calls if stage == "job"]
    assert len(job_payloads) == 1
    assert "SQL preferred" in job_payloads[0] and "Build Python APIs." in job_payloads[0]
    assert '"minimumExperience": 2.0' in job_payloads[0]


def test_jd_file_only_and_pdf_docx_uploads(client):
    response = submit(client, data={"jobDescription": "", "minimumExperience": ""}, files=[
        ("jobDescriptionFile", ("job.txt", b"Python required", "text/plain")),
        ("resumes", ("alice.pdf", pdf_bytes(), "application/pdf")),
        ("resumes", ("alice.docx", docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
    ])
    assert response.status_code == 200, response.text
    assert response.json()["processedCandidates"] == 2


def test_github_is_associated_by_candidate_name_not_upload_order(client, agents):
    profiles = [{"candidateName": "Alice Example", "githubURL": "https://github.com/example/project"}]
    response = submit(client, data={"githubProfiles": json.dumps(profiles)}, files=[
        ("resumes", ("bob.txt", b"Bob Example", "text/plain")),
        ("resumes", ("alice.txt", b"Alice Example", "text/plain")),
    ])
    assert response.status_code == 200, response.text
    alice, bob = response.json()["candidates"]
    assert alice["github"] == "Available" and alice["projects"][0]["score"] == 85
    assert bob["github"] == "Not Provided" and bob["projects"] == []
    matching = [json.loads(payload["github_data"]) for stage, payload in agents.calls if stage == "matching"]
    assert [data["status"] for data in matching] == ["not_provided", "available"]


def test_profile_urls_resolve_to_existing_repository_helper(client, agents, monkeypatch):
    class Response:
        def raise_for_status(self):
            pass

        def json(self):
            return [{"html_url": "https://github.com/example/project", "fork": False, "archived": False}]

    monkeypatch.setattr(orchestrator.requests, "get", lambda *args, **kwargs: Response())
    response = submit(client, data={"githubProfiles": '[{"resumeIndex":0,"githubURL":"https://github.com/example"}]'})
    assert response.status_code == 200
    assert ("github_fetch", {"url": "https://github.com/example/project"}) in agents.calls


def test_duplicate_candidate_names_do_not_share_evidence(client):
    response = submit(client, data={"githubProfiles": '[{"candidateName":"Alice Example","githubURL":"https://github.com/example/project"}]'},
                      files=[("resumes", ("one.txt", b"Alice Example", "text/plain")),
                             ("resumes", ("two.txt", b"Alice Example", "text/plain"))])
    assert response.status_code == 200
    assert all(item["github"] == "Not Provided" and item["warnings"] for item in response.json()["candidates"])


def test_candidate_failures_do_not_erase_successes(client):
    response = submit(client, files=[("resumes", ("alice.txt", b"Alice Example", "text/plain")),
                                     ("resumes", ("bad.pdf", b"bad PDF", "application/pdf")),
                                     ("resumes", ("bad-agent.txt", b"BROKEN_AGENT", "text/plain"))])
    assert response.status_code == 200
    result = response.json()
    assert result["status"] == "partial"
    assert result["processedCandidates"] == 1 and result["totalCandidates"] == 3
    assert [error["stage"] for error in result["errors"]] == ["parsing", "resume"]
    assert "private provider diagnostic" not in response.text


def test_missing_github_evidence_is_nonfatal(client, agents, monkeypatch):
    def unavailable(url):
        raise RuntimeError("private request details")
    monkeypatch.setattr(agents, "github_data", unavailable)
    response = submit(client, data={"githubProfiles": '[{"resumeFilename":"alice.txt","githubURL":"https://github.com/example/project"}]'})
    assert response.status_code == 200
    candidate = response.json()["candidates"][0]
    assert candidate["github"] == "Unavailable" and candidate["warnings"]
    assert "private request details" not in response.text


@pytest.mark.parametrize("data", [
    {"jobRole": " "}, {"jobDescription": ""}, {"minimumExperience": "-1"}, {"minimumExperience": "nan"},
    {"githubProfiles": "not json"}, {"githubProfiles": "{}"},
    {"githubProfiles": '[{"resumeIndex":9,"githubURL":"https://github.com/example"}]'},
    {"githubProfiles": '[{"candidateName":"Alice","githubURL":"https://github.com.evil.test/example"}]'},
    {"githubProfiles": '[{"candidateName":"Alice","githubURL":"http://127.0.0.1/private"}]'},
    {"githubProfiles": '[{"candidateName":"Alice","resumeIndex":0,"githubURL":"https://github.com/example"}]'},
])
def test_invalid_input_is_rejected_before_agents(client, agents, data):
    assert submit(client, data=data).status_code == 422
    assert agents.calls == []


def test_upload_limits_and_empty_batch(client, agents):
    assert submit(client, files=[]).status_code == 422
    assert submit(client, files=[("resumes", ("large.txt", b"a" * (10 * 1024 * 1024 + 1), "text/plain"))]).status_code == 422
    assert submit(client, files=[("resumes", (f"{i}.txt", b"Alice Example", "text/plain")) for i in range(101)]).status_code == 422
    assert agents.calls == []


def test_missing_matching_score_never_becomes_zero(client, agents, monkeypatch):
    original = agents.invoke
    monkeypatch.setattr(agents, "invoke", lambda stage, payload: "No score provided" if stage == "matching" else original(stage, payload))
    response = submit(client)
    assert response.status_code == 502
    assert response.json()["errors"][0]["stage"] == "matching"


def test_cors_uses_explicit_local_origins(client):
    response = client.options("/api/analyze", headers={"Origin": "http://localhost:5500", "Access-Control-Request-Method": "POST"})
    assert response.headers["access-control-allow-origin"] == "http://localhost:5500"
    response = client.options("/api/analyze", headers={"Origin": "https://unrelated.example", "Access-Control-Request-Method": "POST"})
    assert "access-control-allow-origin" not in response.headers
